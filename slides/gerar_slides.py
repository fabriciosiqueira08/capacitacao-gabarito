#!/usr/bin/env python3
"""Gera os decks da capacitação de back-end a partir do deck de front-end do Gabriel Fiuza.

O arquivo do Canva (`template/_Capacitação Front-end.pptx`) exportou tudo achatado: todos os
slides usam o layout `Blank` e o design vive em shapes posicionados absolutamente dentro de cada
slide — herdar o layout não herda nada. Então a técnica aqui é clonar slides inteiros ("doadores")
e trocar só o texto, o que preserva fundo, faixas, logo e cores originais.

    python3 slides/gerar_slides.py            # gera as quatro aulas
    python3 slides/gerar_slides.py --aula 1   # só a Aula 1
"""

import argparse
import copy
import re
from pathlib import Path

import yaml
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

RAIZ = Path(__file__).resolve().parent
TEMPLATE = RAIZ / "template" / "_Capacitação Front-end.pptx"
CONTEUDO = RAIZ / "conteudo"
BUILD = RAIZ / "build"

LARGURA_IN = 20.0  # o deck do Canva é 20" x 11.25"

# Paleta extraída do próprio arquivo do Fiuza.
AZUL_ESCURO = RGBColor(0x0E, 0x27, 0x66)
CINZA_CODIGO = RGBColor(0xE6, 0xEA, 0xF2)
CIANO = RGBColor(0x48, 0xB1, 0xFF)

# Slide doador de cada arquétipo, e em quais shapes desse doador mora o texto.
# (índices conferidos no arquivo original — ver README de slides/)
ARQUETIPOS = {
    "capa": {"doador": 1, "titulo": 1, "corpo": 3, "remover": [2]},
    "objetivos": {"doador": 2, "titulo": 1, "corpo": 2},
    "secao": {"doador": 17, "titulo": 1},
    "titulo": {"doador": 6, "titulo": 1},
    "bullets": {"doador": 22, "titulo": 1, "corpo": 2},
    # Doador 85 é o slide de prática do Fiuza, mas o corpo dele é centralizado e
    # estreito. Usamos o corpo largo do 22 e só o azul do título de prática.
    "pratica": {"doador": 22, "titulo": 1, "corpo": 2, "cor_titulo": CIANO},
    "codigo": {"doador": 22, "titulo": 1, "corpo": 2, "codigo": True},
    "duas_colunas": {
        "doador": 9,
        "titulo": 2,
        "esquerda": 3,
        "direita": 4,
        "direita_titulo": 5,
        "esquerda_titulo": 6,
    },
    "fim": {"doador": 138, "titulo": 1, "corpo": 2},
}


# --------------------------------------------------------------------------- clonagem


def clonar_slide(destino, slide_origem):
    """Copia um slide inteiro (shapes + imagens) do deck de origem para o de destino."""
    novo = destino.slides.add_slide(destino.slide_layouts[6])  # 'Blank'
    for forma in list(novo.placeholders):
        forma._element.getparent().remove(forma._element)

    mapa_rid = {}
    for rid, rel in slide_origem.part.rels.items():
        if rel.reltype == RT.IMAGE:
            mapa_rid[rid] = novo.part.relate_to(rel._target, RT.IMAGE)

    # A capa pinta o fundo em <p:bg>, fora do spTree — sem isso ela sai branca.
    fundo = slide_origem._element.find(qn("p:cSld")).find(qn("p:bg"))
    if fundo is not None:
        novo._element.find(qn("p:cSld")).insert(0, copy.deepcopy(fundo))

    for forma in slide_origem.shapes:
        elemento = copy.deepcopy(forma._element)
        _remapear_rids(elemento, mapa_rid)
        novo.shapes._spTree.append(elemento)

    return novo


def _remapear_rids(elemento, mapa_rid):
    """As imagens copiadas ainda apontam para os rIds do slide de origem."""
    for no in elemento.iter():
        for atributo, valor in list(no.attrib.items()):
            if atributo.endswith("}embed") or atributo.endswith("}link"):
                if valor in mapa_rid:
                    no.attrib[atributo] = mapa_rid[valor]


# --------------------------------------------------------------------------- texto


def _modelos_de_paragrafo(text_frame):
    """Um parágrafo-modelo por nível de indentação, para clonar preservando a formatação."""
    modelos = {}
    for paragrafo in text_frame.paragraphs:
        nivel = paragrafo.level
        if nivel not in modelos and paragrafo.runs:
            modelos[nivel] = copy.deepcopy(paragrafo._p)
    if not modelos:
        raise ValueError("shape doador sem parágrafo com texto")
    return modelos


def escrever(shape, linhas, ajustar=False, centralizar=False, cor=None):
    """Reescreve um shape de texto mantendo fonte, cor, bullet e espaçamento do doador.

    `linhas` é uma lista de strings; uma linha iniciada por "- " vira sub-bullet.
    Com `ajustar`, a fonte encolhe até o texto caber na caixa do doador.
    """
    escala = _escala_que_cabe(shape, linhas) if ajustar else 1.0
    # Só o corpo interessa: título encolhido continua enorme.
    if escala < LIMITE_LEGIVEL and not centralizar:
        AVISOS.append(f"letra em {escala:.0%} — considere dividir: {linhas[0][:50]!r}")
    text_frame = shape.text_frame
    modelos = _modelos_de_paragrafo(text_frame)
    nivel_base = min(modelos)
    nivel_sub = max(modelos)

    corpo = text_frame._txBody
    for paragrafo in corpo.findall(qn("a:p")):
        corpo.remove(paragrafo)

    for linha in linhas:
        nivel = nivel_sub if linha.startswith("- ") else nivel_base
        texto = linha[2:].strip() if linha.startswith("- ") else linha
        novo = copy.deepcopy(modelos[nivel])
        corpo.append(novo)
        _definir_texto_do_paragrafo(novo, texto)

    # ⚬ (U+26AC) do doador vira tofu em quem não tem a fonte do Canva.
    for marcador in corpo.iter(qn("a:buChar")):
        if marcador.get("char") == "⚬":
            marcador.set("char", "◦")

    if escala != 1.0:
        _escalar_fontes(corpo, escala)
    if centralizar:
        _centralizar(shape)
    if cor is not None:
        _pintar(corpo, cor)


def _definir_texto_do_paragrafo(paragrafo, texto):
    execucoes = paragrafo.findall(qn("a:r"))
    for extra in execucoes[1:]:
        paragrafo.remove(extra)
    execucoes[0].find(qn("a:t")).text = texto


def _escalar_fontes(corpo, escala):
    for no in corpo.iter():
        if no.tag == qn("a:rPr") and no.get("sz"):
            no.set("sz", str(max(900, int(int(no.get("sz")) * escala))))
        elif no.tag == qn("a:spcPts") and no.get("val"):
            no.set("val", str(max(1000, int(int(no.get("val")) * escala))))


def _pintar(corpo, cor):
    for no in corpo.iter(qn("a:srgbClr")):
        no.set("val", str(cor))


def _centralizar(shape):
    """Ocupa a largura útil do slide e centraliza — evita título comprido estourando a caixa."""
    shape.left = Inches(1.0)
    shape.width = Inches(LARGURA_IN - 2.0)
    for paragrafo in shape.text_frame.paragraphs:
        paragrafo.alignment = PP_ALIGN.CENTER


ALTURA_IN = 11.25
LARGURA_GLIFO = 0.52  # largura média de caractere em ems, para a fonte do deck
ALTURA_LINHA = 1.6  # entrelinha em ems
# O slide mais denso do Fiuza usa 28pt; o nosso doador de corpo tem 52,63pt.
# 28/52,63 ≈ 0,53 — abaixo disso é mais apertado do que ele jamais escreveu.
LIMITE_LEGIVEL = 0.5
AVISOS = []


def _escala_que_cabe(shape, linhas):
    """Maior fator de fonte (≤ 1) em que as linhas cabem na caixa do doador.

    Sem isso, um slide com mais bullets que o doador transborda a área útil —
    o Canva reajusta sozinho, o python-pptx não.
    """
    fonte_pt = _fonte_do_doador(shape)
    largura = shape.width / 914400
    altura = min(shape.height / 914400, ALTURA_IN - shape.top / 914400 - 0.4)

    escala = 1.0
    while escala > 0.35:
        pt = fonte_pt * escala
        por_linha = max(12, int(largura * 72 / (pt * LARGURA_GLIFO)))
        visuais = sum(max(1, -(-len(linha) // por_linha)) for linha in linhas)
        if visuais * pt * ALTURA_LINHA / 72 <= altura:
            return escala
        escala -= 0.05
    return 0.35


def _fonte_do_doador(shape):
    for paragrafo in shape.text_frame.paragraphs:
        for execucao in paragrafo.runs:
            if execucao.font.size:
                return execucao.font.size.pt
    return 40.0


# --------------------------------------------------------------------------- código


def caixa_de_codigo(slide, codigo, topo=Inches(2.6)):
    """Bloco de código: retângulo escuro + texto monoespaçado. Fica editável no Canva."""
    linhas = codigo.rstrip().split("\n")
    largura_max = max(len(linha) for linha in linhas)

    corpo_pt = min(32.0, 1250.0 / max(largura_max, 28), 460.0 / max(len(linhas), 6))
    altura = Inches(len(linhas) * corpo_pt * 1.45 / 72 + 0.9)
    # Fator generoso de propósito: a máquina de quem abrir pode substituir a
    # monoespaçada por uma mais larga, e aí a última linha quebraria.
    largura = Inches(min(18.6, largura_max * corpo_pt * 0.72 / 72 + 1.1))
    esquerda = Inches((LARGURA_IN - largura.inches) / 2)

    fundo = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, esquerda, topo, largura, altura)
    fundo.fill.solid()
    fundo.fill.fore_color.rgb = AZUL_ESCURO
    fundo.line.fill.background()
    fundo.shadow.inherit = False
    fundo.adjustments[0] = 0.04

    caixa = slide.shapes.add_textbox(
        esquerda + Inches(0.5), topo + Inches(0.4), largura - Inches(1.0), altura - Inches(0.8)
    )
    text_frame = caixa.text_frame
    text_frame.word_wrap = False
    for indice, linha in enumerate(linhas):
        paragrafo = text_frame.paragraphs[0] if indice == 0 else text_frame.add_paragraph()
        execucao = paragrafo.add_run()
        execucao.text = linha or " "
        execucao.font.name = "Consolas"
        execucao.font.size = Pt(corpo_pt)
        execucao.font.color.rgb = CINZA_CODIGO
        paragrafo.line_spacing = 1.25
    return fundo


# --------------------------------------------------------------------------- montagem


def _remover(formas, indices):
    """Remove pelo elemento XML — os índices continuam válidos na lista capturada."""
    for indice in indices:
        elemento = formas[indice]._element
        elemento.getparent().remove(elemento)


def montar_slide(destino, origem, spec):
    tipo = spec["tipo"]
    arquetipo = ARQUETIPOS[tipo]
    slide = clonar_slide(destino, origem.slides[arquetipo["doador"] - 1])
    formas = list(slide.shapes)

    def _shape(indice):
        return formas[indice]

    if "titulo" in arquetipo and spec.get("titulo"):
        escrever(
            _shape(arquetipo["titulo"]),
            spec["titulo"].split("\n"),
            ajustar=True,
            centralizar=True,
            cor=arquetipo.get("cor_titulo"),
        )

    if tipo == "duas_colunas":
        escrever(_shape(arquetipo["esquerda_titulo"]), [spec["esquerda_titulo"]])
        escrever(_shape(arquetipo["direita_titulo"]), [spec["direita_titulo"]])
        for lado in ("esquerda", "direita"):
            escrever(_shape(arquetipo[lado]), spec[lado], ajustar=True)
    elif arquetipo.get("codigo"):
        _remover(formas, [arquetipo["corpo"]])
        caixa_de_codigo(slide, spec["codigo"])
    elif "corpo" in arquetipo:
        linhas = _linhas_do_corpo(spec)
        if linhas:
            escrever(_shape(arquetipo["corpo"]), linhas, ajustar=True)
        else:
            _remover(formas, [arquetipo["corpo"]])

    if arquetipo.get("remover"):
        _remover(formas, arquetipo["remover"])

    if spec.get("notas"):
        slide.notes_slide.notes_text_frame.text = spec["notas"].strip()

    return slide


def _linhas_do_corpo(spec):
    if spec.get("bullets"):
        return [str(item) for item in spec["bullets"]]
    if spec.get("corpo"):
        return spec["corpo"].rstrip().split("\n")
    return []


def gerar(numero_aula):
    caminho = CONTEUDO / f"aula-{numero_aula:02d}.yml"
    roteiro = yaml.safe_load(caminho.read_text(encoding="utf-8"))

    origem = Presentation(TEMPLATE)
    destino = Presentation(TEMPLATE)
    _esvaziar(destino)

    for spec in roteiro["slides"]:
        montar_slide(destino, origem, spec)

    BUILD.mkdir(exist_ok=True)
    saida = BUILD / f"aula-{numero_aula:02d}-{_slug(roteiro['titulo'])}.pptx"
    destino.save(saida)
    return saida, len(roteiro["slides"])


def _esvaziar(prs):
    lista = prs.slides._sldIdLst
    for slide_id in list(lista):
        prs.part.drop_rel(slide_id.rId)
        lista.remove(slide_id)


def _slug(texto):
    texto = re.sub(r"[^\w\s-]", "", texto.lower(), flags=re.UNICODE)
    for de, para in zip("áàâãéêíóôõúç", "aaaaeeiooouc"):
        texto = texto.replace(de, para)
    return re.sub(r"[\s_]+", "-", texto).strip("-")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--aula", type=int, choices=[1, 2, 3, 4], help="gera só uma aula")
    argumentos = parser.parse_args()

    aulas = [argumentos.aula] if argumentos.aula else [1, 2, 3, 4]
    for numero in aulas:
        if not (CONTEUDO / f"aula-{numero:02d}.yml").exists():
            print(f"aula-{numero:02d}.yml ainda não existe — pulando")
            continue

        AVISOS.clear()
        saida, total = gerar(numero)
        print(f"{saida.relative_to(RAIZ.parent)} — {total} slides")
        for aviso in AVISOS:
            print(f"  ! {aviso}")


if __name__ == "__main__":
    main()
